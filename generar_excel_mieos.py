# -*- coding: utf-8 -*-
from io import BytesIO
from openpyxl import Workbook, load_workbook

CABECERAS = [
    "CUENTA", "FECHA FACTURA", "Nº FACTURA", "CONCEPTO", "CIF/NIF", "NOMBRE",
    "DOMICILIO", "LOCALIDAD", "C.P.", "BASE", "% IVA/IGIC", "CUOTA IVA/IGIC",
    "% RECARGO", "CUOTA RECARGO", "% RETENCION", "CUOTA RETENCION",
    "TIPO DE OPERACIÓN", "CLAVE DE GASTOS DIV", "TOTAL FACTURA",
]
FORMATOS_NUMERO = {10: "0.00", 11: "0.00", 12: "0.00", 13: "0.00", 14: "0.00", 16: "0.00", 19: "0.00"}
COLUMNA_TEXTO_FACTURA = 3
ANCHOS_COLUMNA = {
    "A": 10.57, "B": 49.42, "C": 15.28, "D": 29.42, "E": 11.71, "F": 36.42,
    "G": 36.42, "H": 11.71, "I": 7.0, "J": 8.14, "K": 12.85, "L": 17.57,
    "M": 11.71, "N": 16.43, "O": 14.0, "P": 18.71, "Q": 21.14, "R": 23.42, "S": 16.43,
}
TOLERANCIA_CUADRE = 0.02


def _escribe_celda(ws, fila, columna, valor):
    """Escribe un valor; un texto que empiece por = + - @ se guarda como texto, nunca como fórmula de Excel."""
    celda = ws.cell(row=fila, column=columna, value=valor)
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
        celda.data_type = "s"
    return celda


def _valida_cuadre(factura):
    suma = 0.0
    for d in factura["desgloses"]:
        suma += d.get("base", 0) or 0
        suma += d.get("cuota_iva", 0) or 0
        suma += d.get("cuota_recargo", 0) or 0
        suma -= d.get("cuota_retencion", 0) or 0
    return abs(suma - factura["total_factura"]) <= TOLERANCIA_CUADRE


def _clasifica(factura, cif_empresa):
    if factura.get("tipo_conocido") in ("COMPRAS", "VENTAS"):
        return factura["tipo_conocido"]
    cif_emisor = (factura["emisor"].get("cif_nif") or "").strip().upper()
    cif_dest = (factura["destinatario"].get("cif_nif") or "").strip().upper()
    cif_empresa = cif_empresa.strip().upper()
    if cif_emisor == cif_empresa:
        return "VENTAS"   # si el emisor es la propia empresa, es una venta (aunque la lectura repita su CIF como comprador)
    if cif_dest == cif_empresa:
        return "COMPRAS"
    return None


def _concepto(num_factura, nombre_contraparte, max_len=24):
    return f"{num_factura} {(nombre_contraparte or '')[:max_len]}".strip()


def _escribe_hoja(ws, cif_empresa, nombre_empresa, periodo, filas):
    ws["A1"] = "EMPRESA:"
    ws["B1"] = f"{cif_empresa} - {nombre_empresa}"
    ws["A2"] = "PERIODO:"
    ws["B2"] = periodo
    for col_idx, titulo in enumerate(CABECERAS, start=1):
        ws.cell(row=3, column=col_idx, value=titulo)
        if col_idx in FORMATOS_NUMERO:
            ws.cell(row=3, column=col_idx).number_format = FORMATOS_NUMERO[col_idx]
        if col_idx == COLUMNA_TEXTO_FACTURA:
            ws.cell(row=3, column=col_idx).number_format = "@"
    fila_actual = 4
    for fila in filas:
        for col_idx, valor in enumerate(fila, start=1):
            celda = _escribe_celda(ws, fila_actual, col_idx, valor)
            if col_idx in FORMATOS_NUMERO:
                celda.number_format = FORMATOS_NUMERO[col_idx]
            if col_idx == COLUMNA_TEXTO_FACTURA:
                celda.number_format = "@"
        fila_actual += 1
    for letra, ancho in ANCHOS_COLUMNA.items():
        ws.column_dimensions[letra].width = ancho


def _procesa_una_factura(factura, cif_empresa):
    if factura.get("_motivo_revision"):   # el motivo concreto (p. ej. «las dos lecturas no coinciden en: total») va antes que el genérico
        return None, factura["_motivo_revision"]
    if factura.get("confianza") == "baja":
        return None, "el agente no está seguro de haber leído bien algún dato (letra pequeña, borrosa o cortada)"
    tipo = _clasifica(factura, cif_empresa)
    if tipo is None:
        return None, "el CIF de la factura no es el de esta empresa (ni como emisor ni como destinatario)"
    if factura.get("_motivo_revision"):
        return None, factura["_motivo_revision"]
    if not _valida_cuadre(factura):
        return None, "las bases y el IVA no suman el total de la factura"
    from validaciones import problemas_factura
    problemas = problemas_factura(factura)
    if problemas:
        return None, problemas[0]

    contraparte = factura["emisor"] if tipo == "COMPRAS" else factura["destinatario"]
    concepto = _concepto(factura["num_factura"], contraparte.get("nombre", ""))

    def _r(v):
        return round(v, 2) if isinstance(v, (int, float)) else v

    filas = []
    for d in factura["desgloses"]:
        filas.append([
            "C", factura["fecha_factura"], factura["num_factura"], concepto,
            contraparte.get("cif_nif", ""), contraparte.get("nombre", ""),
            contraparte.get("domicilio", "") if tipo == "COMPRAS" else "",
            contraparte.get("localidad", "") if tipo == "COMPRAS" else "",
            contraparte.get("cp", "") if tipo == "COMPRAS" else "",
            _r(d.get("base", 0)), _r(d.get("pct_iva", 0)), _r(d.get("cuota_iva", 0)),
            _r(d.get("pct_recargo") or ""), _r(d.get("cuota_recargo") or ""),
            _r(d.get("pct_retencion") or ""), _r(d.get("cuota_retencion") or ""),
            "", "", _r(factura["total_factura"]),
        ])
    return tipo, filas


def generar_excel_mieos(cif_empresa, nombre_empresa, periodo, facturas, ruta_salida):
    filas_compras, filas_ventas, revision = [], [], []
    for factura in facturas:
        tipo, resultado = _procesa_una_factura(factura, cif_empresa)
        if tipo is None:
            revision.append({"factura": factura, "motivo": resultado})
            continue
        (filas_compras if tipo == "COMPRAS" else filas_ventas).extend(resultado)

    wb = Workbook()
    wb.active.title = "COMPRAS"
    _escribe_hoja(wb.active, cif_empresa, nombre_empresa, periodo, filas_compras)
    ws_v = wb.create_sheet("VENTAS")
    _escribe_hoja(ws_v, cif_empresa, nombre_empresa, periodo, filas_ventas)
    wb.save(ruta_salida)
    return {"ok_compras": len(filas_compras), "ok_ventas": len(filas_ventas), "revision": revision}


def anadir_factura_a_excel(contenido_existente, cif_empresa, nombre_empresa, periodo, factura):
    tipo, resultado = _procesa_una_factura(factura, cif_empresa)
    if tipo is None:
        return contenido_existente, {"ok": False, "motivo": resultado}
    filas_nuevas = resultado

    if contenido_existente:
        wb = load_workbook(BytesIO(contenido_existente))
    else:
        wb = Workbook()
        wb.active.title = "COMPRAS"
        _escribe_hoja(wb.active, cif_empresa, nombre_empresa, periodo, [])
        ws_v = wb.create_sheet("VENTAS")
        _escribe_hoja(ws_v, cif_empresa, nombre_empresa, periodo, [])

    ws = wb["COMPRAS"] if tipo == "COMPRAS" else wb["VENTAS"]
    fila_destino = ws.max_row + 1
    for fila in filas_nuevas:
        for col_idx, valor in enumerate(fila, start=1):
            celda = _escribe_celda(ws, fila_destino, col_idx, valor)
            if col_idx in FORMATOS_NUMERO:
                celda.number_format = FORMATOS_NUMERO[col_idx]
            if col_idx == COLUMNA_TEXTO_FACTURA:
                celda.number_format = "@"
        fila_destino += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), {"ok": True, "tipo": tipo}


def _clave_orden(fila, posicion):
    """Orden de las filas de una hoja: por fecha de factura y luego por número (AC-9 antes que AC-10); a igualdad, como estaban."""
    import re
    from datetime import datetime
    fecha = fila[1]
    try:
        f = datetime.strptime(str(fecha).strip(), "%d/%m/%Y") if isinstance(fecha, str) else datetime(fecha.year, fecha.month, fecha.day)
    except (ValueError, AttributeError, TypeError):
        f = datetime.max
    num = str(fila[2] or "")
    trozos = [(0, int(t), "") if t.isdigit() else (1, 0, t.lower()) for t in re.findall(r"\d+|\D+", num)]
    return (f, trozos, posicion)


def ordenar_filas(contenido):
    """Devuelve el Excel con las filas de COMPRAS y VENTAS ordenadas por fecha y número de factura."""
    wb = load_workbook(BytesIO(contenido))
    for nombre in ("COMPRAS", "VENTAS"):
        if nombre not in wb.sheetnames:
            continue
        ws = wb[nombre]
        filas = []
        for r in range(4, ws.max_row + 1):
            valores = [ws.cell(row=r, column=c).value for c in range(1, len(CABECERAS) + 1)]
            if any(v not in (None, "") for v in valores):
                filas.append(valores)
        ordenadas = [f for _, f in sorted(((_clave_orden(f, i), f) for i, f in enumerate(filas)), key=lambda x: x[0])]
        if ordenadas == filas:
            continue
        for i, valores in enumerate(ordenadas):
            for c, v in enumerate(valores, start=1):
                celda = _escribe_celda(ws, 4 + i, c, v)
                if c in FORMATOS_NUMERO:
                    celda.number_format = FORMATOS_NUMERO[c]
                if c == COLUMNA_TEXTO_FACTURA:
                    celda.number_format = "@"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_contiene(contenido, num_factura, total):
    """True si alguna fila del Excel ya tiene esa factura (mismo número y total): evita duplicar al reintentar tras un corte."""
    wb = load_workbook(BytesIO(contenido), read_only=True)
    num = str(num_factura or "").strip()
    for nombre in ("COMPRAS", "VENTAS"):
        if nombre not in wb.sheetnames:
            continue
        for fila in wb[nombre].iter_rows(min_row=4, values_only=True):
            if fila and len(fila) >= 19 and str(fila[2] or "").strip() == num:
                try:
                    if abs(float(fila[18] or 0) - float(total or 0)) <= 0.01:
                        return True
                except (TypeError, ValueError):
                    pass
    return False
