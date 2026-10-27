# -*- coding: utf-8 -*-
"""
«Clientes»: la lista de clientes de la gestoría, con tres cosas que hacer en cada uno:
  - elegir cómo se agrupan sus facturas en el Excel (como el resto, mensual, trimestral o un solo Excel con todas las fechas),
  - quitarlo de la lista (no se borra ninguna carpeta ni archivo),
  - y, arriba, añadir un cliente a mano (para no depender de que el agente lo deduzca de una factura).
"""

import csv
import io
import logging
import re
import threading

from flask import Blueprint, render_template_string, request, session, redirect, url_for, flash

from almacenamiento_tenant import storage_para_tenant
from basededatos import (alta_automatica_activa, guardar_alta_automatica, reactivar_cliente, listar_clientes, guardar_periodicidad_cliente, quitar_cliente, obtener_periodicidad,
                         obtener_estado_conexiones)
from estilo_web import pagina
from onboarding import alta_empresa_cliente
from panel_web import login_requerido
from validar_cif import comprobar as comprobar_cif

log = logging.getLogger("clientes_web")
bp_clientes = Blueprint("clientes", __name__)

OPCIONES = (("", "Como el resto"), ("mensual", "Un Excel al mes"), ("trimestral", "Un Excel al trimestre"),
            ("unico", "Un solo Excel, con todas las fechas"))
NOMBRE_POR_DEFECTO = {"mensual": "un Excel al mes", "trimestral": "un Excel al trimestre", "unico": "un solo Excel con todas las fechas"}

PLANTILLA = pagina("Clientes", """
<main class="pagina">
  <h1 style="font-size:34px;">Clientes</h1>
  <p class="tenue" style="margin-top:10px;">Aquí están los clientes de su gestoría. Por defecto el agente crea {{ por_defecto }} para cada uno; puede cambiarlo cliente a cliente.
  «Un solo Excel» recoge todas las facturas aunque sean de meses distintos: la fecha de cada línea la puede ajustar usted en el Excel al contabilizarla.</p>
  {% with m = get_flashed_messages() %}{% if m %}<p class="mensaje" role="status" style="margin-top:16px;">{{ m[0] }}</p>{% endif %}{% endwith %}

  {% if trabajo %}
    <p class="mensaje" role="status" style="margin-top:16px;">
      {% if trabajo.terminado %}Importación terminada: {{ trabajo.nuevos }} añadido(s){% if trabajo.existentes %}, {{ trabajo.existentes }} ya estaban{% endif %}{% if trabajo.errores %}, {{ trabajo.errores|length }} con problemas{% endif %}.
      {% else %}Importando clientes: {{ trabajo.hechos }} de {{ trabajo.total }}. Tarda unos segundos por cliente; puede seguir usando el panel.<meta http-equiv="refresh" content="6">{% endif %}
    </p>
    {% if trabajo.errores %}<ul class="tenue" style="margin:8px 0 0 18px;">{% for e in trabajo.errores[:15] %}<li>{{ e }}</li>{% endfor %}</ul>{% endif %}
  {% endif %}

  <form method="post" action="{{ url_for('clientes.nuevo') }}" class="estado-linea" style="margin-top:20px; gap:8px; flex-wrap:wrap;">
    <input name="nombre" placeholder="Nombre del cliente" aria-label="Nombre del cliente" required style="max-width:260px;">
    <input name="cif" placeholder="CIF/NIF" aria-label="CIF o NIF" required style="max-width:150px;">
    <button class="btn" type="submit">Añadir cliente</button>
  </form>

  <details style="margin-top:14px;">
    <summary class="tenue" style="cursor:pointer;">Importar varios clientes a la vez</summary>
    <form method="post" action="{{ url_for('clientes.importar') }}" enctype="multipart/form-data" style="margin-top:10px; display:grid; gap:8px; max-width:560px;">
      <p class="tenue" style="margin:0;">Pegue una línea por cliente (CIF y nombre, en cualquier orden, separados por tabulador, «;» o coma), o suba un Excel o CSV con esas columnas.</p>
      <textarea name="lista" rows="6" placeholder="B12345674;Mi Cliente SL&#10;12345678Z;Nombre Apellido" aria-label="Lista de clientes"></textarea>
      <input type="file" name="archivo" accept=".xlsx,.csv,.txt" aria-label="Archivo de clientes">
      <button class="btn" type="submit" style="justify-self:start;">Importar clientes</button>
    </form>
  </details>

  <form method="post" action="{{ url_for('clientes.ajustes') }}" style="margin-top:14px;">
    <label style="display:flex; gap:10px; align-items:flex-start; max-width:640px; cursor:pointer;">
      <input type="checkbox" name="alta_automatica" value="1" {% if alta_automatica %}checked{% endif %} onchange="this.form.submit()" style="flex:none; width:18px; height:18px; min-width:0; padding:0; margin:3px 0 0;">
      <span>Dar de alta clientes nuevos automáticamente con su primera factura.
        <span class="tenue">Si lo desmarca, solo se archivan facturas de los clientes de esta lista; las de empresas que no estén irán a «Sin identificar» para que decida usted.</span></span>
    </label>
  </form>

  {% if clientes %}
  <input id="buscar" type="search" placeholder="Buscar un cliente…" aria-label="Buscar un cliente" style="margin-top:18px; max-width:320px;">
  <div class="tabla" style="margin-top:12px;"><table class="apilable" id="tabla-clientes">
    <tr><th>Cliente</th><th>Excel</th><th></th></tr>
    {% for c in clientes %}
    <tr data-busca="{{ (c.nombre ~ ' ' ~ c.cif_nif)|lower }}">
      <td data-label="Cliente">{{ c.nombre }}<span class="sec">{{ c.cif_nif }}</span></td>
      <td data-label="Excel">
        <form method="post" action="{{ url_for('clientes.periodicidad', empresa_id=c.id) }}">
          <select name="periodicidad" aria-label="Excel de {{ c.nombre }}" onchange="this.form.submit()">
            {% for valor, texto in opciones %}<option value="{{ valor }}" {% if c.periodicidad == valor %}selected{% endif %}>{{ texto }}</option>{% endfor %}
          </select>
        </form>
      </td>
      <td data-label=""><form method="post" action="{{ url_for('clientes.quitar', empresa_id=c.id) }}"><button class="btn discreto" type="submit"
        onclick="return confirm('¿Quitar a {{ c.nombre|replace("'", "") }} de la lista? No se borra ninguna carpeta ni archivo. Si le llega una factura como comprador, irá a «Sin identificar».');">Quitar cliente</button></form></td>
    </tr>
    {% endfor %}
  </table></div>
  <script>
    document.getElementById("buscar").addEventListener("input", function () {
      var q = this.value.toLowerCase();
      document.querySelectorAll("#tabla-clientes tr[data-busca]").forEach(function (f) { f.style.display = f.dataset.busca.indexOf(q) >= 0 ? "" : "none"; });
    });
  </script>
  {% else %}
  <div class="vacio" style="margin-top:22px;"><strong>Todavía no hay clientes.</strong>Se darán de alta solos con la primera factura, o puede añadir uno arriba.</div>
  {% endif %}
</main>
""", activa="clientes")


_TRABAJOS = {}   # progreso de la importación de cada gestoría (un solo proceso de servidor: basta con memoria)
RE_CIF = re.compile(r"(?<![A-Za-z0-9])([XYZxyz]\d{7}[A-Za-z]|\d{8}[A-Za-z]|[A-Za-z]\d{7}[0-9A-Za-z])(?![A-Za-z0-9])")
MAX_IMPORTAR = 400


def leer_clientes(texto):
    """De un texto (una línea por cliente, CIF y nombre en cualquier orden) saca [(nombre, cif)] y los avisos de líneas no entendidas."""
    clientes, avisos, vistos = [], [], set()
    for linea in (texto or "").splitlines():
        original = linea.strip()
        if not original:
            continue
        m = RE_CIF.search(original)
        if not m:
            avisos.append(f"«{original[:50]}»: no se encuentra un CIF")
            continue
        nombre = re.sub(r"[;\t,|]+", " ", original[:m.start()] + " " + original[m.end():])
        nombre = re.sub(r"\s+", " ", nombre).strip(" -–:.")
        cif, error = comprobar_cif(m.group(1))
        if error or not nombre:
            avisos.append(f"«{original[:50]}»: " + ("el CIF no es válido (revise el dígito o la letra final)" if error else "falta el nombre"))
        elif cif not in vistos:
            vistos.add(cif)
            clientes.append((nombre, cif))
    return clientes, avisos


def _texto_de_archivo(archivo):
    datos = archivo.read(5 * 1024 * 1024)
    if (archivo.filename or "").lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(datos), read_only=True, data_only=True).active
        return "\n".join(";".join("" if v is None else str(v) for v in fila) for fila in hoja.iter_rows(values_only=True))
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return ""


def _importar_en_segundo_plano(tenant_id, clientes):
    trabajo = _TRABAJOS[tenant_id]
    try:
        storage = storage_para_tenant(tenant_id)
        for nombre, cif in clientes:
            try:
                if reactivar_cliente(tenant_id, cif):
                    trabajo["nuevos"] += 1
                elif alta_empresa_cliente(storage, tenant_id, nombre, cif).get("nueva"):
                    trabajo["nuevos"] += 1
                else:
                    trabajo["existentes"] += 1
            except Exception:
                log.exception("No se pudo importar el cliente %s", cif)
                trabajo["errores"].append(f"{nombre} ({cif}): no se pudo crear")
            trabajo["hechos"] += 1
    except Exception:
        log.exception("Fallo importando clientes")
        trabajo["errores"].append("Se interrumpió la importación. Compruebe que el almacenamiento sigue conectado y repítala: los ya añadidos no se duplican.")
    finally:
        trabajo["terminado"] = True


@bp_clientes.route("/clientes/importar", methods=["POST"])
@login_requerido
def importar():
    tenant_id = session["tenant_id"]
    actual = _TRABAJOS.get(tenant_id)
    if actual and not actual["terminado"]:
        flash("Ya hay una importación en marcha; espere a que termine.")
        return redirect(url_for("clientes.lista"))
    texto = request.form.get("lista", "")
    archivo = request.files.get("archivo")
    try:
        if archivo and archivo.filename:
            texto += "\n" + _texto_de_archivo(archivo)
    except Exception:
        log.exception("No se pudo leer el archivo de clientes")
        flash("No se ha podido leer el archivo. Use un Excel (.xlsx), un CSV o pegue la lista.")
        return redirect(url_for("clientes.lista"))
    clientes, avisos = leer_clientes(texto)
    if not clientes:
        flash("No he encontrado ningún cliente en lo que ha enviado. Cada línea necesita un CIF y un nombre.")
        return redirect(url_for("clientes.lista"))
    if len(clientes) > MAX_IMPORTAR:
        flash(f"Son demasiados de golpe (máximo {MAX_IMPORTAR}). Divida la lista.")
        return redirect(url_for("clientes.lista"))
    if not storage_para_tenant(tenant_id):
        flash("Conecte primero su almacenamiento en Conexiones.")
        return redirect(url_for("clientes.lista"))
    _TRABAJOS[tenant_id] = {"total": len(clientes), "hechos": 0, "nuevos": 0, "existentes": 0, "errores": list(avisos), "terminado": False}
    threading.Thread(target=_importar_en_segundo_plano, args=(tenant_id, clientes), daemon=True).start()
    return redirect(url_for("clientes.lista"))


@bp_clientes.route("/clientes/ajustes", methods=["POST"])
@login_requerido
def ajustes():
    activa = request.form.get("alta_automatica") == "1"
    guardar_alta_automatica(session["tenant_id"], activa)
    flash("Los clientes nuevos se darán de alta solos con su primera factura." if activa else
          "A partir de ahora solo se archivan facturas de los clientes de esta lista; las demás irán a «Sin identificar».")
    return redirect(url_for("clientes.lista"))


@bp_clientes.route("/clientes")
@login_requerido
def lista():
    tenant_id = session["tenant_id"]
    return render_template_string(PLANTILLA, clientes=listar_clientes(tenant_id), opciones=OPCIONES,
                                  trabajo=_TRABAJOS.get(tenant_id), alta_automatica=alta_automatica_activa(tenant_id),
                                  por_defecto=NOMBRE_POR_DEFECTO.get(obtener_periodicidad(tenant_id), "un Excel al mes"),
                                  gestoria_nombre=obtener_estado_conexiones(tenant_id)["nombre"])


@bp_clientes.route("/clientes/<empresa_id>/periodicidad", methods=["POST"])
@login_requerido
def periodicidad(empresa_id):
    if guardar_periodicidad_cliente(session["tenant_id"], empresa_id, request.form.get("periodicidad", "")):
        flash("Guardado. Vale para las facturas nuevas; los Excel que ya existen no se tocan.")
    else:
        flash("No se ha podido guardar el cambio.")
    return redirect(url_for("clientes.lista"))


@bp_clientes.route("/clientes/<empresa_id>/quitar", methods=["POST"])
@login_requerido
def quitar(empresa_id):
    nombre = quitar_cliente(session["tenant_id"], empresa_id)
    flash(f"{nombre} ya no está en la lista. Sus carpetas y archivos no se han tocado." if nombre else "Ese cliente ya no estaba en la lista.")
    return redirect(url_for("clientes.lista"))


@bp_clientes.route("/clientes/nuevo", methods=["POST"])
@login_requerido
def nuevo():
    tenant_id = session["tenant_id"]
    nombre, (cif, error) = request.form.get("nombre", "").strip(), comprobar_cif(request.form.get("cif", ""))
    if not nombre or error:
        flash(error or "Escriba el nombre del cliente.")
        return redirect(url_for("clientes.lista"))
    try:
        storage = storage_para_tenant(tenant_id)
        if not storage:
            flash("Conecte primero su almacenamiento en Conexiones.")
            return redirect(url_for("clientes.lista"))
        if reactivar_cliente(tenant_id, cif):
            flash(f"{nombre} vuelve a estar en la lista.")
            return redirect(url_for("clientes.lista"))
        r = alta_empresa_cliente(storage, tenant_id, nombre, cif)
        flash(f"{nombre} añadido." if r.get("nueva") else "Ese CIF ya estaba en la lista (o es un cliente que quitó antes).")
    except Exception:
        log.exception("No se pudo añadir el cliente")
        flash("No se ha podido crear el cliente. Compruebe que el almacenamiento sigue conectado.")
    return redirect(url_for("clientes.lista"))
